import re
import json
import time
import os
import psutil

from pathlib import Path

from parsers.experience_parser import parse_experience
from parsers.resume_section_classifier import segment_resume
from parsers.skill_extraction_engine import extract_skills
from parsers.education_parser import parse_education
from parsers.education_relevance import calculate_education_relevance

from pypdf import PdfReader
from docx import Document

from utils.logger import get_logger


logger = get_logger(__name__)


# ============================================================
# PRE-COMPILED REGEX PATTERNS
# ============================================================

CANDIDATE_ID_PATTERN = re.compile(
    r"resume[\_-]?(\d+)",
    re.IGNORECASE
)

CONTROL_CHARACTER_PATTERN = re.compile(
    r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]"
)

BULLET_PATTERN = re.compile(
    r"[•●▪▫◦‣⁃]"
)

EXCESSIVE_NEWLINES_PATTERN = re.compile(
    r"\n{3,}"
)

MULTIPLE_SPACES_PATTERN = re.compile(
    r"[ \t]+"
)


# ============================================================
# CANDIDATE ID
# ============================================================

def get_candidate_id(file_path):
    """
    Generate a stable candidate ID from the resume filename.
    """

    file_name = Path(file_path).stem.lower()

    match = CANDIDATE_ID_PATTERN.search(file_name)

    if match:
        number = int(match.group(1))
        return f"CAND-{number:03d}"

    # Fallback for test or custom filenames.
    return "CAND-001"


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_pdf_text(file_path):
    """
    Extract text from a text-based PDF resume safely.

    Handles:
    - Corrupted PDF files
    - Invalid PDF files
    - Empty PDFs
    - PDFs without extractable text

    Raises a clear ValueError when extraction fails.
    """

    try:
        reader = PdfReader(str(file_path))

        pages = []

        for page in reader.pages:
            page_text = page.extract_text()

            if page_text:
                pages.append(page_text)

    except Exception as exc:
        logger.exception(
            "Failed to read PDF resume: %s",
            file_path
        )

        raise ValueError(
            f"Unable to read PDF resume "
            f"'{Path(file_path).name}'. "
            "The file may be corrupt or invalid."
        ) from exc

    text = "\n".join(pages)

    if not text.strip():
        raise ValueError(
            "No text found in PDF. "
            "It may be a scanned document requiring OCR."
        )

    logger.info(
        "Text extracted from PDF: %s",
        file_path
    )

    return text


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_docx_text(file_path):
    """
    Extract paragraphs and table content from a DOCX resume.

    Handles:
    - Corrupted DOCX files
    - Invalid Word documents
    - Empty documents
    - Documents without extractable text
    """

    try:
        document = Document(str(file_path))

        parts = []

        # Extract non-empty paragraphs.
        for paragraph in document.paragraphs:
            paragraph_text = paragraph.text.strip()

            if paragraph_text:
                parts.append(paragraph_text)

        # Extract non-empty table rows.
        for table in document.tables:
            for row in table.rows:

                row_text = []

                for cell in row.cells:
                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:
                    parts.append(" | ".join(row_text))

    except Exception as exc:
        logger.exception(
            "Failed to read DOCX resume: %s",
            file_path
        )

        raise ValueError(
            f"Unable to read DOCX resume "
            f"'{Path(file_path).name}'. "
            "The file may be corrupt or invalid."
        ) from exc

    text = "\n".join(parts)

    if not text.strip():
        raise ValueError(
            "No text found in DOCX resume."
        )

    logger.info(
        "Text extracted from DOCX: %s",
        file_path
    )

    return text


# ============================================================
# RESUME TEXT CLEANING
# ============================================================

def clean_resume_text(text):
    """
    Clean whitespace, bullet symbols,
    and unwanted control characters.
    """

    if not text or not text.strip():
        raise ValueError(
            "Resume text cannot be empty."
        )

    # Normalize line endings.
    text = text.replace(
        "\r\n",
        "\n"
    ).replace(
        "\r",
        "\n"
    )

    # Remove unwanted control characters.
    text = CONTROL_CHARACTER_PATTERN.sub(
        "",
        text
    )

    # Standardize common bullet symbols.
    text = BULLET_PATTERN.sub(
        "-",
        text
    )

    # Clean individual lines.
    lines = []

    for line in text.split("\n"):
        cleaned_line = line.strip()
        lines.append(cleaned_line)

    # Remove excessive blank lines.
    text = "\n".join(lines)

    text = EXCESSIVE_NEWLINES_PATTERN.sub(
        "\n\n",
        text
    )

    # Normalize repeated spaces and tabs.
    text = MULTIPLE_SPACES_PATTERN.sub(
        " ",
        text
    )

    return text.strip()


# ============================================================
# BASIC RESUME PARSING
# ============================================================

def parse_resume_text(text):
    """
    Validate resume text and return cleaned resume content.
    """

    if not text or not text.strip():
        raise ValueError(
            "Resume text cannot be empty."
        )

    cleaned_text = clean_resume_text(text)

    logger.info(
        "Resume text received and cleaned."
    )

    return {
        "raw_text": text.strip(),
        "cleaned_text": cleaned_text,
        "status": "processed"
    }


# ============================================================
# COMPLETE RESUME EXTRACTION PIPELINE
# ============================================================

def extract_resume(
    file_path,
    output_dir="data/extracted_resumes"
):
    """
    Extract, clean, segment, extract skills and experience,
    and save structured resume data.

    Performance information includes:
    - Text extraction time
    - Text cleaning time
    - Section detection time
    - Total processing time
    - Memory before processing
    - Memory after processing
    - Difference in memory usage
    """

    # --------------------------------------------------------
    # Performance timer
    # --------------------------------------------------------

    total_start = time.perf_counter()

    # --------------------------------------------------------
    # Memory monitoring
    # --------------------------------------------------------

    process = psutil.Process(os.getpid())

    memory_before = (
        process.memory_info().rss
        / (1024 * 1024)
    )

    # --------------------------------------------------------
    # Validate input file
    # --------------------------------------------------------

    file_path = Path(file_path)

    if not file_path.is_file():
        raise FileNotFoundError(
            f"Resume file not found: {file_path}"
        )

    extension = file_path.suffix.lower()

    if extension not in {".pdf", ".docx"}:
        raise ValueError(
            "Unsupported file type. Use PDF or DOCX."
        )

    # --------------------------------------------------------
    # Step 1: Generate candidate ID
    # --------------------------------------------------------

    candidate_id = get_candidate_id(file_path)

    # --------------------------------------------------------
    # Step 2: Extract raw text
    # --------------------------------------------------------

    extraction_start = time.perf_counter()

    if extension == ".pdf":
        raw_text = extract_pdf_text(file_path)

    else:
        raw_text = extract_docx_text(file_path)

    extraction_time = (
        time.perf_counter()
        - extraction_start
    )

    # --------------------------------------------------------
    # Step 3: Clean extracted text
    # --------------------------------------------------------

    cleaning_start = time.perf_counter()

    result = parse_resume_text(raw_text)

    cleaning_time = (
        time.perf_counter()
        - cleaning_start
    )

    # --------------------------------------------------------
    # Step 4: Detect resume sections
    # --------------------------------------------------------

    section_start = time.perf_counter()

    sections = segment_resume(
        result["cleaned_text"]
    )

    section_time = (
        time.perf_counter()
        - section_start
    )

    # --------------------------------------------------------
    # Step 5: Extract skills
    # --------------------------------------------------------

    skills_text = sections.get(
        "skills",
        ""
    )

    if skills_text and skills_text.strip():
        skills = extract_skills(skills_text)

    else:
        skills = []

    # --------------------------------------------------------
    # Step 6: Extract experience
    # --------------------------------------------------------

    experience_text = sections.get(
        "work_experience",
        ""
    )

    if experience_text and experience_text.strip():

        experience = parse_experience(
            experience_text,
            candidate_id=candidate_id
        )

    else:

        experience = {
            "candidate_id": candidate_id,
            "experiences": [],
            "total_experience_months": 0,
            "gaps": [],
            "overlaps": []
        }

    # --------------------------------------------------------
    # Step 7: Add structured sections
    # --------------------------------------------------------

    result["sections"] = sections

    # --------------------------------------------------------
    # Step 8: Add extracted skills
    # --------------------------------------------------------

    result["extracted_skills"] = skills

    # --------------------------------------------------------
    # Step 9: Extract education and certifications
    # --------------------------------------------------------

    education_data = parse_education(
        result["cleaned_text"],
        candidate_id=candidate_id
    )

    result["education"] = education_data["education"]

    result["certifications"] = (
        education_data["certifications"]
    )

    result["education_relevance"] = {
        "score": 0.0,
        "matched_education": [],
        "matched_certifications": []
    }

    # --------------------------------------------------------
    # Step 10: Add structured experience
    # --------------------------------------------------------

    result["experience"] = experience

    # --------------------------------------------------------
    # Step 11: Create output directory
    # --------------------------------------------------------

    output_path = Path(output_dir)

    output_path.mkdir(
        parents=True,
        exist_ok=True
    )

    # --------------------------------------------------------
    # Step 12: Create output filename
    # --------------------------------------------------------

    output_file = (
        output_path
        / f"{file_path.stem}.json"
    )

    # --------------------------------------------------------
    # Step 13: Measure processing time and memory
    # --------------------------------------------------------

    total_time = (
        time.perf_counter()
        - total_start
    )

    memory_after = (
        process.memory_info().rss
        / (1024 * 1024)
    )

    memory_used = (
        memory_after
        - memory_before
    )

    # --------------------------------------------------------
    # Step 14: Add performance information
    # --------------------------------------------------------

    result["performance"] = {
        "text_extraction_seconds": round(
            extraction_time,
            6
        ),

        "text_cleaning_seconds": round(
            cleaning_time,
            6
        ),

        "section_detection_seconds": round(
            section_time,
            6
        ),

        "total_processing_seconds": round(
            total_time,
            6
        ),

        "memory_before_mb": round(
            memory_before,
            3
        ),

        "memory_after_mb": round(
            memory_after,
            3
        ),

        "memory_used_mb": round(
            memory_used,
            3
        )
    }

    # --------------------------------------------------------
    # Step 15: Prepare output data
    # --------------------------------------------------------

    output_data = {
        "candidate_id": candidate_id,
        "source_file": file_path.name,
        **result
    }

    # --------------------------------------------------------
    # Step 16: Save structured result
    # --------------------------------------------------------

    try:
        with open(
            output_file,
            "w",
            encoding="utf-8"
        ) as file:

            json.dump(
                output_data,
                file,
                indent=4,
                ensure_ascii=False
            )

    except OSError:
        logger.exception(
            "Failed to save structured resume output: %s",
            output_file
        )
        raise

    # --------------------------------------------------------
    # Step 17: Logging
    # --------------------------------------------------------

    logger.info(
        "Structured resume output saved: %s",
        output_file
    )

    logger.info(
        "Resume processing completed in %.4f seconds",
        total_time
    )

    logger.info(
        "Memory usage: %.3f MB before, "
        "%.3f MB after, %.3f MB difference",
        memory_before,
        memory_after,
        memory_used
    )

    # --------------------------------------------------------
    # Step 18: Return structured result
    # --------------------------------------------------------

    return {
        "candidate_id": candidate_id,
        **result
    }